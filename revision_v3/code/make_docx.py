#!/usr/bin/env python3
"""Convert the markdown deliverables to DOCX and apply the journal manuscript styling
(Times New Roman 12 pt, double line spacing, 1 inch margins, embedded figures)."""
import os, subprocess, sys
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
MS = os.path.join(ROOT, "manuscript")
OUT = os.path.join(ROOT, "deliverable")
os.makedirs(OUT, exist_ok=True)


def style(path, double=True):
    doc = Document(path)
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(1)
        s.top_margin = s.bottom_margin = Inches(1)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    body = ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text")
    for p in doc.paragraphs:
        if p.style.name in body:
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
            p.paragraph_format.space_after = Pt(0)
        for r in p.runs:
            if r.font.name is None:
                r.font.name = "Times New Roman"
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                    for r in p.runs:
                        r.font.size = Pt(9)
    doc.save(path)
    print("styled", os.path.basename(path))


def convert(md, out, resource=None):
    cmd = ["pandoc", md, "-o", out, "--wrap=none", "-s",
           "--resource-path", resource or os.path.dirname(md)]
    subprocess.run(cmd, check=True, cwd=os.path.dirname(md))
    style(out)


def main():
    convert(os.path.join(MS, "main.md"), os.path.join(OUT, "main_clean.docx"), MS)
    convert(os.path.join(MS, "supplementary.md"),
            os.path.join(OUT, "supplementary.docx"), MS)
    convert(os.path.join(OUT, "Response_to_Reviewers.md"),
            os.path.join(OUT, "Response_to_Reviewers.docx"), OUT)
    convert(os.path.join(OUT, "cover_letter.md"),
            os.path.join(OUT, "cover_letter.docx"), OUT)
    print("docx done")


if __name__ == "__main__":
    main()
